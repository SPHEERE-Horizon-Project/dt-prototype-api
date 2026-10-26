"""Quantity selection, physical units, dependencies and unchanged assessment totals."""
from copy import deepcopy
import json
from pathlib import Path
from types import SimpleNamespace

import numpy as np
import pandas as pd
import pytest

from dt_prototype.assessment.catalogue import load_json, validate_catalogue
from dt_prototype.assessment.contract import AssessmentInputError, REQUIRED_COLUMNS
from dt_prototype.assessment.quantities import (resolve_quantities, run_quantities,
    ProjectQuantityContext, check_assessment_quantities, load_quantity_inputs)
from dt_prototype.assessment.workflow import run_assessment
from tests.dynamic_rc.conftest import make_box_building, make_constant_weather

ROOT = Path(__file__).resolve().parents[2]
CAT = load_json(ROOT/'configs/asset_catalogue.json')


def asset(key='pv', cid='pv_system', **updates):
    row = dict(asset_id=key,catalogue_id=cid,scope_action='installed',source='explicit test input')
    row.update(updates)
    return row


def inputs(*rows, default='auto', bid='001'):
    return dict(schema_version='1.0',default_mode=default,
                buildings=[dict(assessment_id='test',building_id=bid,assets=list(rows))])


def driver(bid, key):
    values = dict(roof_area=100,opaque_wall_area=200,ground_area=80,glazed_area=20,
                  floor_area=300,heating_design=10,cooling_design=8,dhw_daily_volume=400,
                  dhw_daily_energy=20)
    return values[key], 'independent test driver'


def test_auto_factor_and_direct_input_have_distinct_meanings():
    auto = asset(rule_id='fixed_estimate',parameters={'estimate':10},auto_factor=1.2,direct_quantity=999)
    result = resolve_quantities(inputs(auto),CAT).iloc[0]
    assert (result.base_quantity,result.sizing_factor,result.assessed_quantity) == (10,1.2,12)
    auto.update(quantity_mode='direct',direct_quantity=7,auto_factor=-1,rule_id='inactive_invalid_rule')
    result = resolve_quantities(inputs(auto),CAT,lambda *_:pytest.fail('direct mode evaluated estimate')).iloc[0]
    assert (result.base_quantity,result.sizing_factor,result.assessed_quantity) == (7,1,7)
    assert pd.isna(result.estimated_quantity)


def test_direct_zero_and_missing_do_not_fall_back():
    data = inputs(asset(direct_quantity=0),default='direct')
    result = resolve_quantities(data,CAT).iloc[0]
    assert result.assessed_quantity == 0 and result.status == 'complete'
    data['buildings'][0]['assets'][0].pop('direct_quantity')
    result = resolve_quantities(data,CAT,driver).iloc[0]
    assert result.status == 'missing' and pd.isna(result.assessed_quantity)
    assert 'no auto fallback' in result.diagnostic


def test_defaults_and_per_asset_overrides_preserve_building_id_text():
    data = inputs(asset(direct_quantity=3),asset('tank','water_tank',quantity_mode='auto',auto_factor=.5),default='direct')
    data['buildings'].append(dict(assessment_id='test',building_id='1',default_mode='auto',
                                 assets=[asset(rule_id='fixed_estimate',parameters={'estimate':2})]))
    result = resolve_quantities(data,CAT,driver)
    assert result.building_id.tolist() == ['001','001','1']
    assert result.assessed_quantity.tolist() == [3,200,2]


def test_reference_dependency_uses_resolved_direct_quantity_and_is_order_independent():
    data = inputs(asset('battery','battery_system',rule_id='asset_multiple',
                        parameters={'asset_id':'pv','rate':1.5,'rate_unit':'kWh/kW'},auto_factor=1.1),
                  asset(quantity_mode='direct',direct_quantity=10))
    result = resolve_quantities(data,CAT).set_index('asset_id')
    assert result.loc['battery','base_quantity'] == 15
    assert result.loc['battery','assessed_quantity'] == 16.5
    assert result.loc['pv','assessed_quantity'] == 10


def test_dependency_missing_is_reported_and_cycles_are_errors():
    rows = [asset('battery','battery_system',rule_id='asset_multiple',
                  parameters={'asset_id':'pv','rate':2,'rate_unit':'kWh/kW'}),
            asset(quantity_mode='direct')]
    result = resolve_quantities(inputs(*rows),CAT)
    assert result.status.tolist() == ['missing','missing']
    rows[1].update(quantity_mode='auto',rule_id='asset_multiple',
                   parameters={'asset_id':'battery','rate':1,'rate_unit':'kW/kWh'})
    with pytest.raises(AssessmentInputError,match='circular'):
        resolve_quantities(inputs(*rows),CAT)
    rows.pop()
    with pytest.raises(AssessmentInputError,match='unknown referenced'):
        resolve_quantities(inputs(*rows),CAT)


@pytest.mark.parametrize('rule,params,cid,expected',[
    ('opaque_wall_area',{},'facade_insulation',200),
    ('glazed_area',{},'glazing_system',20),
    ('roof_fraction',{'coverage':.25},'solar_thermal_system',25),
    ('pv_roof_capacity',{'coverage':.5,'module_kW_m2':.2},'pv_system',10),
    ('dhw_recovery',{'recovery_hours':5},'dhw_heat_pump',4),
    ('dhw_daily_volume',{},'water_tank',400),
    ('heating_design',{},'heat_pump_system',10),
    ('cooling_design',{},'cooling_system',8),
])
def test_dimensioned_estimate_rules(rule,params,cid,expected):
    row = resolve_quantities(inputs(asset(cid=cid,rule_id=rule,parameters=params)),CAT,driver).iloc[0]
    assert row.assessed_quantity == expected


@pytest.mark.parametrize('rule,params,cid',[
    ('pv_roof_capacity',{'coverage':1.1,'module_kW_m2':.2},'pv_system'),
    ('dhw_recovery',{'recovery_hours':0},'dhw_heat_pump'),
    ('dhw_recovery',{'recovery_hours':25},'dhw_heat_pump'),
    ('fixed_estimate',{'estimate':True},'pv_system'),
    ('fixed_estimate',{'estimate':float('inf')},'pv_system'),
    ('fixed_estimate',{'estimate':-1},'pv_system'),
    ('fixed_estimate',{'estimate':1,'typo':2},'pv_system'),
    ('roof_area',{},'pv_system'),
])
def test_invalid_rule_inputs_and_dimensions_are_rejected(rule,params,cid):
    with pytest.raises(AssessmentInputError):
        resolve_quantities(inputs(asset(cid=cid,rule_id=rule,parameters=params)),CAT,driver)


@pytest.mark.parametrize('value',[True,-1,float('nan'),float('inf'),'10'])
def test_direct_quantity_numeric_domain(value):
    with pytest.raises(AssessmentInputError):
        resolve_quantities(inputs(asset(direct_quantity=value),default='direct'),CAT)


@pytest.mark.parametrize('value',[True,0,-1,float('inf')])
def test_auto_factor_numeric_domain(value):
    with pytest.raises(AssessmentInputError):
        resolve_quantities(inputs(asset(rule_id='fixed_estimate',parameters={'estimate':1},auto_factor=value)),CAT)


def test_overflow_and_incompatible_reference_units_fail():
    with pytest.raises(AssessmentInputError,match='finite'):
        resolve_quantities(inputs(asset(rule_id='fixed_estimate',parameters={'estimate':1e300},auto_factor=1e300)),CAT)
    with pytest.raises(AssessmentInputError,match='rate_unit'):
        resolve_quantities(inputs(asset(quantity_mode='direct',direct_quantity=10),
            asset('battery','battery_system',rule_id='asset_multiple',parameters={'asset_id':'pv','rate':2,'rate_unit':'h'})),CAT)


def test_sum_rule_keeps_shared_generator_capacity_explicit():
    result = resolve_quantities(inputs(asset('space',quantity_mode='direct',direct_quantity=10),
        asset('water',quantity_mode='direct',direct_quantity=2),
        asset('shared',rule_id='asset_sum',parameters={'asset_ids':['space','water']})),CAT)
    assert result.iloc[2].assessed_quantity == 12
    with pytest.raises(AssessmentInputError,match='distinct'):
        resolve_quantities(inputs(asset(rule_id='asset_sum',parameters={'asset_ids':['pv','pv']})),CAT)


def test_bacs_capacity_basis_requires_an_explicit_definition():
    data = inputs(asset(cid='bacs',direct_quantity=10),default='direct')
    assert resolve_quantities(data,CAT).iloc[0].status == 'missing'
    data['buildings'][0]['assets'][0]['quantity_basis'] = 'controlled heating thermal capacity, kW'
    assert resolve_quantities(data,CAT).iloc[0].assessed_quantity == 10


def test_missing_auto_parameters_and_context_are_not_zero():
    result = resolve_quantities(inputs(asset()),CAT,driver).iloc[0]
    assert result.status == 'missing' and 'missing parameters' in result.diagnostic
    result = resolve_quantities(inputs(asset(cid='roof_insulation')),CAT).iloc[0]
    assert result.status == 'missing' and 'context required' in result.diagnostic


def test_invalid_schema_identity_and_unit_are_rejected():
    for data in [inputs(asset(),asset()),inputs(asset(quantity_unit='W')),
                 inputs(asset(catalogue_id='unknown')),inputs(asset(quantity_mode='hybrid')),
                 inputs(asset(source=''))]:
        with pytest.raises(AssessmentInputError):
            resolve_quantities(data,CAT,driver)
    data = inputs(asset())
    data['buildings'].append(deepcopy(data['buildings'][0]))
    with pytest.raises(AssessmentInputError,match='duplicate assessment/building'):
        resolve_quantities(data,CAT,driver)


def test_catalogue_units_and_duplicate_keys_are_checked(tmp_path):
    assert {a['quantity_unit'] for a in validate_catalogue(CAT).values()} == {'m2','kW','kWh','L'}
    cat = deepcopy(CAT)
    cat['assets'].append(deepcopy(cat['assets'][0]))
    with pytest.raises(AssessmentInputError,match='duplicate catalogue_id'):
        validate_catalogue(cat)
    file = tmp_path/'duplicate.json'
    file.write_text('{"schema_version":"1.0","schema_version":"2.0"}')
    with pytest.raises(AssessmentInputError,match='duplicate JSON'):
        load_json(file)


def test_static_project_drivers_keep_inventory_separate_from_heat_loss():
    building = make_box_building(ach=.5)
    weather = make_constant_weather(t_ext=0)
    ctx = ProjectQuantityContext.__new__(ProjectQuantityContext)
    ctx.project = SimpleNamespace(configuration={'operating':{}},config_path=Path('test.json'))
    ctx.district = SimpleNamespace(weather=weather,buildings=[building])
    ctx.cache = {}
    before = ctx._prepare(building)
    # U_wall=.8, U_window=2.8, roof=.9, ground=1.2*.7.
    ua = 240*(.85*.8+.15*2.8) + 100*.9 + 100*1.2*.7
    expected = (ua + 600*1.2/3600*.5*1006)*20/1000
    assert before['heating_design'] == pytest.approx(expected)
    assert before['opaque_wall_area'] == 240*.85
    building.geometry.zones[0].surfaces[0].boundary_condition = 'adiabatic'
    after = ctx._prepare(building)
    assert after['opaque_wall_area'] == before['opaque_wall_area']
    assert after['heating_design'] < before['heating_design']


def test_run_exports_missing_status_and_never_overwrites(tmp_path):
    cat, source = tmp_path/'cat.json',tmp_path/'in.json'
    cat.write_text(json.dumps(CAT))
    source.write_text(json.dumps(inputs(asset(),default='direct')))
    output = run_quantities(source,cat,tmp_path/'out')
    assert json.loads((output/'manifest.json').read_text())['status'] == 'incomplete'
    assert pd.read_csv(output/'resolved_asset_quantities.csv').status.iloc[0] == 'missing'
    with pytest.raises(FileExistsError):
        run_quantities(source,cat,output)


def test_optional_assessment_join_checks_quantities_without_rescaling_results(tmp_path):
    from tests.assessment.test_assessment import _row
    quantity = resolve_quantities(inputs(asset('hp-system','heat_pump_system',
        rule_id='fixed_estimate',parameters={'estimate':10},auto_factor=1.2),bid='B1'),CAT)
    quantity['assessment_id'] = 'retrofit-a'
    qfile = tmp_path/'quantity.csv'
    quantity.to_csv(qfile,index=False)
    data = pd.DataFrame([_row(base_quantity=10,sizing_factor=1.2,assessed_quantity=12,
                              quantity_unit='kW',value=1234)],columns=REQUIRED_COLUMNS)
    file = tmp_path/'results.csv'
    data.to_csv(file,index=False)
    out = run_assessment(file,tmp_path/'assessment',asset_quantities=qfile)
    assert pd.read_csv(out/'building_asset_results.csv').value.iloc[0] == 1234
    data['quantity_unit'] = 'unit'
    data.to_csv(file,index=False)
    with pytest.raises(AssessmentInputError,match='quantity_unit'):
        run_assessment(file,tmp_path/'wrong',asset_quantities=qfile)


def test_fillable_workbook_round_trips_raw_inputs_and_cached_preview():
    import openpyxl
    file = ROOT/'data/examples/asset_quantity_inputs.xlsx'
    original = load_json(ROOT/'configs/asset_quantities_example.json')
    extracted = load_quantity_inputs(file)
    expected = resolve_quantities(original,CAT,driver)
    actual = resolve_quantities(extracted,CAT,driver)
    pd.testing.assert_frame_equal(actual,expected,check_exact=True)
    book = openpyxl.load_workbook(file,read_only=True,data_only=True)
    assert book.sheetnames == ['Quantities','Catalogue']
    assert book['Quantities']['B5'].value == '1'
    assert all(book['Quantities'].cell(r,15).value == 'PASS' for r in range(11,23))
    assert all(book['Quantities'].cell(r,4).value is None for r in range(11,27))
    assert not [(s.title,c.coordinate) for s in book for row in s for c in row if c.data_type == 'e']
    book.close()


@pytest.mark.parametrize('bad_field', ['formula','parameters','headers'])
def test_workbook_reader_rejects_unsupported_editable_inputs(monkeypatch, bad_field):
    import openpyxl
    from dt_prototype.assessment.quantity_workbook import HEADERS, read_quantity_workbook
    row = [None]*10
    row[0],row[1],row[7],row[8] = 'hp','heat_pump_system','installed','test'
    if bad_field == 'formula':
        row[3] = '=100/2'
    elif bad_field == 'parameters':
        row[6] = '{"estimate":NaN}'
    headers = HEADERS if bad_field != 'headers' else ('wrong',)*10
    class Sheet:
        def __getitem__(self,key):
            return SimpleNamespace(value={'B4':'test','B5':'001','B6':'direct','B7':'test'}[key])
        def cell(self,r,c):
            return SimpleNamespace(value=headers[c-1])
        def iter_rows(self,**kwargs):
            return iter([row])
    class Book:
        sheetnames = ['Quantities']
        def __getitem__(self,key):return Sheet()
        def close(self):pass
    monkeypatch.setattr(openpyxl,'load_workbook',lambda *args,**kwargs:Book())
    with pytest.raises(AssessmentInputError):
        read_quantity_workbook('test.xlsx')
