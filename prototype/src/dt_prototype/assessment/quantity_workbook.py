"""Read editable workbook inputs, never cached formula results as model inputs."""
from pathlib import Path
import json

from .contract import AssessmentInputError

HEADERS = ("asset_id", "catalogue_id", "quantity_mode", "direct_quantity", "auto_factor",
           "rule_id", "parameters", "scope_action", "source", "quantity_basis")


def read_quantity_workbook(path):
    try:
        import openpyxl
    except ImportError as error:
        raise AssessmentInputError("Excel inputs need openpyxl; install dt-prototype[quantity-inputs]") from error
    book = openpyxl.load_workbook(Path(path),read_only=True,data_only=False)
    try:
        if "Quantities" not in book.sheetnames:
            raise AssessmentInputError("quantity workbook requires sheet Quantities")
        sheet = book["Quantities"]
        if tuple(sheet.cell(10,c).value for c in range(1,11)) != HEADERS:
            raise AssessmentInputError("quantity workbook input headers must match template row 10")
        def literal(value, label):
            if isinstance(value,str) and value.startswith("="):
                raise AssessmentInputError(f"{label}: editable inputs must be literal values, not formulas")
            return value
        assessment = literal(sheet["B4"].value,"assessment_id")
        building = literal(sheet["B5"].value,"building_id")
        mode = literal(sheet["B6"].value,"default_mode")
        description = literal(sheet["B7"].value,"description") or "Workbook editable inputs; estimates recalculated in Python"
        assets = []
        for i, values in enumerate(sheet.iter_rows(min_row=11,max_col=10,values_only=True),start=11):
            if all(v is None or v == "" for v in values):
                continue
            row = {key:literal(value,f"row {i} {key}") for key,value in zip(HEADERS,values)
                   if value is not None and value != ""}
            if "parameters" in row:
                try:
                    def reject_constant(value):
                        raise ValueError(f"non-finite JSON number {value}")
                    row["parameters"] = json.loads(row["parameters"],parse_constant=reject_constant)
                except (TypeError,ValueError) as error:
                    raise AssessmentInputError(f"row {i}: parameters must be a JSON object") from error
            assets.append(row)
        return dict(schema_version="1.0",description=description,
                    default_mode=mode,buildings=[dict(assessment_id=assessment,building_id=building,assets=assets)])
    finally:
        book.close()
