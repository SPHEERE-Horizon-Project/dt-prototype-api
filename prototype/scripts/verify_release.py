"""Verify the cleaned release, accepted results, workbooks and local links."""
from __future__ import annotations

import hashlib
import argparse
import json
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET

import pandas as pd
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from dt_prototype.common.project import ProjectInputs
from dt_prototype.energy_carriers import EnergyConfig
from dt_prototype import __version__


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def junit_counts(path: Path) -> dict[str, int]:
    suites = list(ET.parse(path).getroot().iter("testsuite"))
    return {key: sum(int(s.attrib.get(key, 0)) for s in suites)
            for key in ("tests", "failures", "errors", "skipped")}


def check_links() -> int:
    checked = 0
    documents = [ROOT / "README.md", *sorted((ROOT / "docs").glob("*.md")),
                 *sorted((ROOT / "runners").glob("*.md")),
                 ROOT / "src/dt_prototype/monthly_ss_vbdd/README.md",
                 ROOT / "src/dt_prototype/monthly_ss_simplified/model3/README.md",
                 ROOT / "reports/workbooks/workbook_validation.md"]
    for document in documents:
        body = document.read_text(encoding="utf-8-sig")
        for target in re.findall(r"!?\[[^\]]*\]\(([^)]+)\)", body):
            if "://" in target or target.startswith("#"):
                continue
            target = target.strip("<>").split("#")[0]
            if target and not (document.parent / target).resolve().exists():
                raise AssertionError(f"Broken local link in {document}: {target}")
            checked += bool(target)
    return checked


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--energy-results", type=Path, default=ROOT / "reports/shared_energy",
                        help="Completed baseline run whose source hashes must match current code")
    parser.add_argument("--expected-tests", type=int, default=273,
                        help="Expected JUnit records, including unittest subtests (frozen run: 273)")
    parser.add_argument("--output", type=Path,
                        help="Audit destination; defaults to the selected run for a fresh run")
    args = parser.parse_args()
    if args.expected_tests < 1:
        parser.error("--expected-tests must be positive")
    required = ["src", "configs", "data", "runners", "tests", "spreadsheets",
                "reports", "docs", "scripts"]
    for name in required:
        if not (ROOT / name).is_dir():
            raise AssertionError(f"Missing release directory: {name}")
    for name in ("build", "tmp", "tools"):
        if (ROOT / name).exists():
            raise AssertionError(f"Obsolete/intermediate directory remains: {name}")

    result_dir = args.energy_results.resolve()
    manifest = json.loads((result_dir / "manifest.json").read_text(encoding="utf-8"))
    if manifest["status"] != "complete":
        raise AssertionError("Accepted energy run is not complete")
    package = ROOT / "src/dt_prototype"
    # Frozen evidence predates Assessment. Fresh manifests cover the complete
    # package; never exclude newly added boundary code from a source audit.
    has_assessment = any(Path(name).parts[0] == "Assessment" for name in manifest["code_hashes"])
    code_hashes = {str(path.relative_to(package)): sha256(path)
                   for path in package.rglob("*.py")
                   if has_assessment or path.relative_to(package).parts[0] != "Assessment"}
    if code_hashes != manifest["code_hashes"]:
        raise AssertionError("Production source differs from the accepted run")
    project = ProjectInputs.load(ROOT / "configs/example.json")
    if project.hashes() != manifest["input_hashes"]:
        raise AssertionError("Example project inputs differ from the accepted run")
    if EnergyConfig.load(result_dir / "energy_config_resolved.json").sha256 != manifest["energy_config_sha256"]:
        raise AssertionError("Accepted energy configuration hash is inconsistent")

    energy_monthly = pd.read_csv(result_dir / "energy_monthly.csv")
    service_monthly = pd.read_csv(result_dir / "service_monthly.csv")
    if len(energy_monthly) != 612 or energy_monthly.building_id.nunique() != 5:
        raise AssertionError("Unexpected accepted monthly-result coverage")
    if len(service_monthly) != 3060:
        raise AssertionError("Unexpected accepted service-result coverage")
    tests = junit_counts(result_dir / "test_results.xml")
    if tests != {"tests": args.expected_tests, "failures": 0, "errors": 0, "skipped": 0}:
        raise AssertionError(f"Accepted test evidence changed: {tests}")

    workbook_manifest = json.loads((ROOT / "reports/workbooks/workbook_manifest.json").read_text(encoding="utf-8"))
    workbook_summary = json.loads((ROOT / "reports/workbooks/native_excel_summary.json").read_text(encoding="utf-8-sig"))
    if workbook_summary["mismatches"] or workbook_summary["formula_errors"]:
        raise AssertionError("Native Excel evidence contains failures")
    if any(not check["pass"] for check in workbook_summary["checks"]):
        raise AssertionError("Native Excel perturbation/check evidence contains failures")
    workbook_statuses = {}
    for key, details in workbook_manifest["artifacts"].items():
        path = ROOT / "spreadsheets" / details["path"]
        if sha256(path) != details["sha256"]:
            raise AssertionError(f"Workbook hash differs from accepted evidence: {path.name}")
        formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
        if formulas._external_links:
            raise AssertionError(f"External workbook link in {path.name}")
        if "Energy_Check" not in formulas.sheetnames:
            raise AssertionError(f"Missing Energy_Check in {path.name}")
        formulas.close()
        values = openpyxl.load_workbook(path, read_only=True, data_only=True)
        states = [values["Energy_Check"].cell(row, 4).value
                  for row in range(7, values["Energy_Check"].max_row + 1)]
        states = [state for state in states if state is not None]
        if not states or any(state != "PASS" for state in states):
            raise AssertionError(f"Cached baseline check is not all PASS in {path.name}")
        workbook_statuses[key] = {"rows": len(states), "sha256": sha256(path)}
        values.close()

    checked_links = check_links()
    outcome = {
        "status": "passed",
        "energy_result_directory": str(result_dir),
        "version": __version__,
        "accepted_energy_cases": len(manifest["cases"]),
        "accepted_energy_months": len(energy_monthly),
        "accepted_service_rows": len(service_monthly),
        "accepted_tests": tests,
        "native_excel_comparisons": workbook_summary["comparisons"],
        "native_excel_checks": len(workbook_summary["checks"]),
        "workbooks": workbook_statuses,
        "checked_local_links": checked_links,
        "production_code_hashes": code_hashes,
    }
    frozen = (ROOT / "reports/shared_energy").resolve()
    destination = args.output or (ROOT / "reports/release_verification.json" if result_dir == frozen
                                 else result_dir / "release_verification.json")
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(outcome, indent=2), encoding="utf-8")
    print(json.dumps({key: value for key, value in outcome.items()
                      if key != "production_code_hashes"}, indent=2))


if __name__ == "__main__":
    main()
